#!/usr/bin/env python3
import os
import re
import sys
import argparse
from typing import Dict, List, Optional, Tuple
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment
from datetime import datetime

REQUIRED_SHEETS = ["NEA", "DAC", "EMS", "WIL"]

def norm(s: str) -> str:
    s = str(s)
    s = s.strip().lower()
    s = s.replace("\u00a0", " ")
    s = s.replace(".", "")
    s = s.replace(",", " ")
    s = re.sub(r"\s+", " ", s)
    return s

PZN_HEADERS = {"#", "pzn"}
BEZ_HEADERS = {"bez", "bezeichnung", "artikel", "artikelname", "name", "produkt"}
HERST_HEADERS = {"herst", "hersteller", "lieferant"}

def find_columns(df: pd.DataFrame) -> Tuple[Optional[str], Optional[str], Optional[str], Dict[str, str]]:
    """Return (pzn_col, bez_col, herst_col, debugmap) where debugmap maps normalized->original."""
    colmap = {norm(c): c for c in df.columns}
    dbg = {k: v for k, v in colmap.items()}
    pzn = next((colmap[k] for k in colmap if k in PZN_HEADERS), None)
    bez = next((colmap[k] for k in colmap if k in BEZ_HEADERS), None)
    herst = next((colmap[k] for k in colmap if k in HERST_HEADERS), None)
    if pzn is None:
        pzn = next((col for k, col in colmap.items() if "#" in k or "pzn" in k), None)
    if bez is None:
        bez = next((col for k, col in colmap.items() if "bez" in k or "bezeich" in k or "artik" in k), None)
    if herst is None:
        herst = next((col for k, col in colmap.items() if "herst" in k or "herstell" in k or "liefer" in k), None)
    return pzn, bez, herst, dbg

def clean_pzn(val) -> Optional[str]:
    if pd.isna(val):
        return None
    s = str(val).strip()
    digits = re.sub(r"\D+", "", s)
    if digits == "":
        return None
    return digits

def autosize_ws(ws):
    widths = {}
    for row in ws.iter_rows(values_only=True):
        for i, v in enumerate(row, start=1):
            l = len(str(v)) if v is not None else 0
            widths[i] = max(widths.get(i, 0), l)
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), 60)

def write_report(wb, report_rows: List[Dict]):
    if "Vereinigungsmenge_Report" in wb.sheetnames:
        ws = wb["Vereinigungsmenge_Report"]
        wb.remove(ws)
    ws = wb.create_sheet("Vereinigungsmenge_Report", 1)
    cols = ["Zeit", "Eingabe-Datei", "Blatt", "Rows_in", "PZN_col", "Bez_col", "Herst_col", "Rows_valid", "Rows_after_dedupe", "Hinweis"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in report_rows:
        ws.append([r.get(k, "") for k in cols])
    autosize_ws(ws)

def build_vereinigungsmenge(input_path: str,
                            inplace: bool = False,
                            output_path: Optional[str] = None,
                            pad_pzn8: bool = False,
                            sort_by_pzn: bool = True) -> str:
    if not os.path.isfile(input_path):
        raise FileNotFoundError(f"Datei nicht gefunden: {input_path}")

    xl = pd.ExcelFile(input_path)
    frames = []
    per_sheet_info = []
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for sh in REQUIRED_SHEETS:
        if sh not in xl.sheet_names:
            per_sheet_info.append({
                "Zeit": now, "Eingabe-Datei": os.path.basename(input_path), "Blatt": sh,
                "Rows_in": 0, "PZN_col": "", "Bez_col": "", "Herst_col": "",
                "Rows_valid": 0, "Rows_after_dedupe": "", "Hinweis": "Blatt fehlt – übersprungen"
            })
            continue

        df = pd.read_excel(input_path, sheet_name=sh, dtype=str)
        rows_in = len(df.index)
        pzn_col, bez_col, herst_col, dbg = find_columns(df)

        hint = []
        if pzn_col is None: hint.append("PZN/# nicht gefunden")
        if bez_col is None: hint.append("Bez. nicht gefunden")
        if herst_col is None: hint.append("Herst. nicht gefunden")

        if pzn_col is None or bez_col is None or herst_col is None:
            per_sheet_info.append({
                "Zeit": now, "Eingabe-Datei": os.path.basename(input_path), "Blatt": sh,
                "Rows_in": rows_in, "PZN_col": pzn_col or "", "Bez_col": bez_col or "", "Herst_col": herst_col or "",
                "Rows_valid": 0, "Rows_after_dedupe": "", "Hinweis": "; ".join(hint) if hint else "Ungültige Spalten"
            })
            continue

        tmp = pd.DataFrame({
            "PZN": df[pzn_col].map(clean_pzn),
            "Bez.": df[bez_col].astype(str).str.strip(),
            "Herst.": df[herst_col].astype(str).str.strip()
        })

        tmp = tmp[tmp["PZN"].notna() & (tmp["PZN"] != "")]
        if pad_pzn8:
            tmp["PZN"] = tmp["PZN"].apply(lambda s: s.zfill(8))

        tmp["Quelle"] = sh
        frames.append(tmp)

        per_sheet_info.append({
            "Zeit": now, "Eingabe-Datei": os.path.basename(input_path), "Blatt": sh,
            "Rows_in": rows_in, "PZN_col": pzn_col, "Bez_col": bez_col, "Herst_col": herst_col,
            "Rows_valid": len(tmp.index), "Rows_after_dedupe": "", "Hinweis": ""
        })

    if not frames:
        raise ValueError("Keine verwertbaren Daten gefunden (fehlende Spalten oder leere Blätter).")

    # Vereinigungsmenge bilden (Union) und nach PZN deduplizieren
    all_rows = pd.concat(frames, axis=0, ignore_index=True)
    before = len(all_rows.index)
    all_rows = all_rows.drop_duplicates(subset=["PZN"], keep="first")
    after = len(all_rows.index)

    # optional sort
    if sort_by_pzn:
        try:
            all_rows = all_rows.sort_values(by="PZN", key=lambda s: s.astype(int), ignore_index=True)
        except Exception:
            all_rows = all_rows.sort_values(by="PZN", ignore_index=True)

    # write output file (copy by default)
    if inplace:
        out_path = input_path
    else:
        base, ext = os.path.splitext(input_path)
        out_path = output_path or f"{base}_mit_Vereinigungsmenge{ext}"
        with open(input_path, "rb") as src, open(out_path, "wb") as dst:
            dst.write(src.read())

    wb = load_workbook(out_path)

    if "Vereinigungsmenge" in wb.sheetnames:
        wb.remove(wb["Vereinigungsmenge"])
    ws = wb.create_sheet("Vereinigungsmenge", 0)

    ws.append(["PZN", "Bez.", "Herst.", "Quelle"])
    for c in ws[1]:
        c.font = Font(bold=True)

    for _, row in all_rows.iterrows():
        ws.append([row["PZN"], row["Bez."], row["Herst."], row.get("Quelle", "")])

    ws.freeze_panes = "A2"

    # report
    total_hint = f"Gesamt: {before} Zeilen vor Deduplizierung; {after} nach Deduplizierung; entfernt: {before-after}"
    per_sheet_info.append({
        "Zeit": now, "Eingabe-Datei": os.path.basename(input_path), "Blatt": "TOTAL",
        "Rows_in": "", "PZN_col": "", "Bez_col": "", "Herst_col": "",
        "Rows_valid": "", "Rows_after_dedupe": after, "Hinweis": total_hint
    })

    if "Vereinigungsmenge_Report" in wb.sheetnames:
        wb.remove(wb["Vereinigungsmenge_Report"])
    rep = wb.create_sheet("Vereinigungsmenge_Report", 1)

    cols = ["Zeit", "Eingabe-Datei", "Blatt", "Rows_in", "PZN_col", "Bez_col", "Herst_col", "Rows_valid", "Rows_after_dedupe", "Hinweis"]
    rep.append(cols)
    for c in rep[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    for r in per_sheet_info:
        rep.append([r.get(k, "") for k in cols])

    # autosize
    for wsx in [ws, rep]:
        try:
            widths = {}
            for row in wsx.iter_rows(values_only=True):
                for i, v in enumerate(row, start=1):
                    l = len(str(v)) if v is not None else 0
                    widths[i] = max(widths.get(i, 0), l)
            for i, w in widths.items():
                wsx.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), 60)
        except Exception:
            pass

    wb.save(out_path)
    return out_path

def find_input_in_folder(folder: str) -> str:
    candidates = [os.path.join(folder, f) for f in os.listdir(folder) if f.lower().endswith(".xlsx")]
    if not candidates:
        raise FileNotFoundError(f"Keine .xlsx-Dateien im Ordner: {folder}")
    candidates.sort(key=lambda p: os.path.getmtime(p), reverse=True)
    return candidates[0]

def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Erzeuge/aktualisiere das Blatt 'Vereinigungsmenge' aus NEA/DAC/EMS/WIL (Union, dedupliziert nach PZN)."
    )
    g = parser.add_mutually_exclusive_group(required=True)
    g.add_argument("--input", help="Pfad zur Excel-Datei mit den Blättern NEA, DAC, EMS, WIL.")
    g.add_argument("--folder", help="Ordner mit der Excel-Datei (nimmt die neueste .xlsx).")
    parser.add_argument("--inplace", action="store_true", help="Direkt im Eingabe-Workbook schreiben (keine Kopie).")
    parser.add_argument("--output", help="Optionaler Ausgabepfad (nur wenn --inplace nicht gesetzt ist).")
    parser.add_argument("--pad-pzn8", action="store_true", help="PZN links auf 8 Stellen mit Nullen auffüllen.")
    parser.add_argument("--no-sort", action="store_true", help="Nicht nach PZN sortieren.")
    args = parser.parse_args(argv)

    input_path = args.input or find_input_in_folder(args.folder)
    out_path = build_vereinigungsmenge(
        input_path,
        inplace=args.inplace,
        output_path=args.output,
        pad_pzn8=args.pad_pzn8,
        sort_by_pzn=(not args.no_sort),
    )
    print(out_path)

if __name__ == "__main__":
    sys.exit(main())